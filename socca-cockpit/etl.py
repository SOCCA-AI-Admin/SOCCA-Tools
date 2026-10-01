#!/usr/bin/env python3
"""
SOCCA GROUP Sales Cockpit — ETL
===============================
Liest
  1) die Sales-Arbeitsmappe   (.xlsx, .xlsm oder .xlsb)
     Blaetter: Sales, FTE, Leads, Props, Goals, Locations
  2) den Combit-Export        (.csv oder .xlsx)
und schreibt einen kompakten Langdatensatz (data.json), aus dem das
Dashboard alles rechnet.

    python3 etl.py Sales.xlsx C_AP.xlsx web/data.json

Beide Eingabeformate werden automatisch erkannt. Formelergebnisse werden
aus dem Wert-Cache der Datei gelesen; eine "Werte only"-Kopie ist nicht
noetig, solange die Mappe zuletzt mit Excel gespeichert wurde.

Definitionen (gegen das Blatt "Report" verifiziert):
  Stichtag        Buchungsdatum (Spalte "Date"), nicht Reisedatum
  DB              MargIn = VK - EK
  Buchungen       Anzahl Zeilen
  Teams           Summe der Spalte "Teams" (eine Buchung kann mehrere Teams enthalten)
  Hotel           Spalte 9 im Blatt Sales = WebID im Combit-Export = WebID im Blatt Locations
  Geschaeftsjahr  Juli bis Juni (AP-Logik)
"""
import sys, os, json, csv, datetime, math
from collections import defaultdict

EPOCH = datetime.date(1899, 12, 30)
DAY0 = datetime.date(2019, 1, 1)
MONTH0 = (2019, 1)
LEAD_CUT = (2024, 1)          # ab hier zaehlt der Combit-Export statt der Mappe


# ------------------------------------------------------- Laender/Regionen
# Combit fuehrt ISO-2, das Blatt Sales die deutsche Kurzform. Angeglichen
# wird auf die Schreibweise des Blattes Sales.
ISO_TO_SALES = {'DE': 'D', 'AT': 'A', 'IT': 'I'}

COUNTRY = {
    'D': 'Deutschland', 'A': 'Österreich', 'CH': 'Schweiz', 'I': 'Italien',
    'HR': 'Kroatien', 'CZ': 'Tschechien', 'ES': 'Spanien', 'SI': 'Slowenien',
    'HU': 'Ungarn', 'TR': 'Türkei', 'DK': 'Dänemark', 'PT': 'Portugal',
    'GR': 'Griechenland', 'SK': 'Slowakei', 'CY': 'Zypern', 'NL': 'Niederlande',
    'PL': 'Polen', 'EG': 'Ägypten', 'MT': 'Malta', 'LU': 'Luxemburg',
    'FI': 'Finnland', 'NO': 'Norwegen', 'FR': 'Frankreich', 'SE': 'Schweden',
    'LI': 'Liechtenstein', 'BE': 'Belgien', 'GB': 'Großbritannien',
    'IE': 'Irland', 'US': 'USA', 'AE': 'VAE', 'RO': 'Rumänien',
    'BG': 'Bulgarien', 'RS': 'Serbien', 'BA': 'Bosnien-Herzegowina',
    'ME': 'Montenegro', 'MK': 'Nordmazedonien', 'AL': 'Albanien',
    'EE': 'Estland', 'LV': 'Lettland', 'LT': 'Litauen', 'UA': 'Ukraine',
}

REGION = {
    'D': {'BW': 'Baden-Württemberg', 'BY': 'Bayern', 'BE': 'Berlin',
          'BB': 'Brandenburg', 'HB': 'Bremen', 'HH': 'Hamburg', 'HE': 'Hessen',
          'MV': 'Mecklenburg-Vorpommern', 'NI': 'Niedersachsen',
          'NW': 'Nordrhein-Westfalen', 'RP': 'Rheinland-Pfalz',
          'SL': 'Saarland', 'SN': 'Sachsen', 'ST': 'Sachsen-Anhalt',
          'SH': 'Schleswig-Holstein', 'TH': 'Thüringen'},
    'A': {'W': 'Wien', 'NOE': 'Niederösterreich', 'OOE': 'Oberösterreich',
          'SBG': 'Salzburg', 'T': 'Tirol', 'VBG': 'Vorarlberg',
          'KTN': 'Kärnten', 'STMK': 'Steiermark', 'BGLD': 'Burgenland'},
    'CH': {'ZH': 'Zürich', 'BE': 'Bern', 'BN': 'Bern', 'LU': 'Luzern',
           'UR': 'Uri', 'SZ': 'Schwyz', 'OW': 'Obwalden', 'NW': 'Nidwalden',
           'GL': 'Glarus', 'ZG': 'Zug', 'FR': 'Freiburg', 'SO': 'Solothurn',
           'BS': 'Basel-Stadt', 'BL': 'Basel-Landschaft', 'SH': 'Schaffhausen',
           'AR': 'Appenzell Ausserrhoden', 'AI': 'Appenzell Innerrhoden',
           'SG': 'St. Gallen', 'GR': 'Graubünden', 'AG': 'Aargau',
           'TG': 'Thurgau', 'TI': 'Tessin', 'VD': 'Waadt', 'VS': 'Wallis',
           'NE': 'Neuenburg', 'GE': 'Genf', 'JU': 'Jura'},
}


def norm_country(v):
    """ISO-2 auf die Schreibweise des Blattes Sales bringen."""
    t = (v or '').strip().upper()
    if not t:
        return None
    return ISO_TO_SALES.get(t, t)


def norm_region(v):
    """Umlaute in Regionskuerzeln vereinheitlichen (NOE, OOE)."""
    t = (v or '').strip().upper()
    return t.replace('\u00d6', 'OE').replace('\u00c4', 'AE').replace('\u00dc', 'UE') or None


def region_label(land, code):
    name = REGION.get(land, {}).get(code)
    return f'{code} - {name}' if name else code


# ------------------------------------------------------------------ Helfer
def dayidx(d):   return (d - DAY0).days
def monthidx(d): return (d.year - MONTH0[0]) * 12 + (d.month - MONTH0[1])


def num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return None if (math.isnan(v) or math.isinf(v)) else float(v)
    return None


def numlike(v):
    """Wie num(), akzeptiert zusaetzlich Zahlen als Text (CSV-Spalten)."""
    n = num(v)
    if n is not None:
        return n
    if isinstance(v, str):
        s = v.strip().replace(' ', '')
        if not s:
            return None
        try:
            return float(s.replace('.', '').replace(',', '.') if s.count(',') == 1
                         and s.count('.') > 1 else s.replace(',', '.'))
        except ValueError:
            return None
    return None


def as_date(v):
    """Excel-Serienzahl, datetime oder deutscher Datumstext -> date."""
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    n = num(v)
    if n is not None:
        if not (20000 < n < 60000):
            return None
        try:
            return EPOCH + datetime.timedelta(days=n)
        except Exception:
            return None
    if isinstance(v, str):
        s = v.strip()
        for f in ('%d.%m.%Y %H:%M:%S', '%d.%m.%Y %H:%M', '%d.%m.%Y',
                  '%Y-%m-%d %H:%M:%S', '%Y-%m-%d'):
            try:
                return datetime.datetime.strptime(s, f).date()
            except ValueError:
                pass
    return None


def text(v):
    if isinstance(v, str):
        s = v.strip()
        # #REF!, #VALUE! und Co. kommen je nach Leser als 0x17 oder #REF! an
        return None if (not s or s.startswith('0x') or s.startswith('#')) else s
    return None


def r2(v, nd=2):
    return None if v is None else round(v, nd)


# ----------------------------------------------------------- Mappen-Leser
class Workbook:
    """Einheitlicher Zugriff auf .xlsb und .xlsx/.xlsm.

    sheet(name) liefert Zeilen als dict {Spaltenindex ab 0: Wert}.
    """

    def __init__(self, path):
        self.path = path
        ext = os.path.splitext(path)[1].lower()
        if ext == '.xlsb':
            from pyxlsb import open_workbook
            self.kind = 'xlsb'
            self.wb = open_workbook(path)
            self.names = list(self.wb.sheets)
        elif ext in ('.xlsx', '.xlsm', '.xltx'):
            import openpyxl
            self.kind = 'xlsx'
            self.wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            self.names = list(self.wb.sheetnames)
        else:
            raise SystemExit(f'Unbekanntes Mappenformat: {path}')

    def resolve(self, name):
        """Blattnamen tolerant finden (Gross/Klein, Leerzeichen)."""
        norm = lambda s: s.strip().lower().replace(' ', '')
        for n in self.names:
            if norm(n) == norm(name):
                return n
        raise SystemExit(f'Blatt "{name}" fehlt in {self.path}. '
                         f'Vorhanden: {", ".join(self.names)}')

    def sheet(self, name, limit=None):
        real = self.resolve(name)
        if self.kind == 'xlsb':
            with self.wb.get_sheet(real) as sh:
                for i, row in enumerate(sh.rows()):
                    if limit is not None and i > limit:
                        break
                    yield {c.c: c.v for c in row}
        else:
            ws = self.wb[real]
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if limit is not None and i > limit:
                    break
                yield {j: v for j, v in enumerate(row) if v is not None}


    def cells(self, name, max_row):
        """Zellen mit Kennung, ob eine Formel dahintersteht:
        liefert (Zeile, Spalte, Wert, ist_formel), Zeilen ab 0."""
        real = self.resolve(name)
        if self.kind == 'xlsb':
            from pyxlsb import biff12
            with self.wb.get_sheet(real) as sh:
                sh._reader.seek(sh._data_offset, os.SEEK_SET)
                r = -1
                for item in sh._reader:
                    rec, obj = item[0], item[1]
                    if rec == biff12.ROW:
                        r = obj.r
                        if r > max_row:
                            break
                    elif biff12.BLANK <= rec <= biff12.FORMULA_BOOLERR:
                        v = obj.v
                        if rec == biff12.STRING and sh._stringtable is not None:
                            v = sh._stringtable[v]
                        yield r, obj.c, v, rec >= biff12.FORMULA_STRING
                    elif rec == biff12.SHEETDATA_END:
                        break
        else:
            import openpyxl
            fw = openpyxl.load_workbook(self.path, read_only=True, data_only=False)
            vals = self.wb[real].iter_rows(max_row=max_row + 1, values_only=True)
            forms = fw[real].iter_rows(max_row=max_row + 1, values_only=True)
            for r, (vr, fr) in enumerate(zip(vals, forms)):
                for c, (v, f) in enumerate(zip(vr, fr)):
                    if v is not None or f is not None:
                        yield r, c, v, isinstance(f, str) and f.startswith('=')
            fw.close()


# ---------------------------------------------------------------- Mappe
SALES_COL = dict(Date=2, Arrival=1, Sport=4, Team=5, Reg=6, Region=7, Hotel=9,
                 Land=10, Destination=11, VK=12, EK=13, Margin=15,
                 Teams=24, Pax=27, Nights=28)


def read_workbook(path):
    wb = Workbook(path)
    out = {}

    # --- Sales: eine Zeile je Buchung ------------------------------------
    bookings = []
    for i, d in enumerate(wb.sheet('Sales')):
        if i == 0:
            continue
        bd = as_date(d.get(SALES_COL['Date']))
        if bd is None or not (2015 < bd.year < 2040):
            continue
        hotel = num(d.get(SALES_COL['Hotel']))
        bookings.append(dict(
            d=dayidx(bd),
            team=text(d.get(SALES_COL['Team'])),
            sport=text(d.get(SALES_COL['Sport'])),
            reg=norm_country(text(d.get(SALES_COL['Reg']))),
            region=norm_region(text(d.get(SALES_COL['Region']))),
            land=norm_country(text(d.get(SALES_COL['Land']))),
            hotel=int(hotel) if hotel and 9000 < hotel < 100000 else None,
            name=text(d.get(SALES_COL['Destination'])),
            vk=num(d.get(SALES_COL['VK'])) or 0.0,
            ek=num(d.get(SALES_COL['EK'])) or 0.0,
            marg=num(d.get(SALES_COL['Margin'])) or 0.0,
            teams=num(d.get(SALES_COL['Teams'])) or 0.0,
            pax=num(d.get(SALES_COL['Pax'])) or 0.0,
            nights=num(d.get(SALES_COL['Nights'])) or 0.0,
        ))
    out['bookings'] = bookings

    # --- Monatsreihen je Team: FTE, Leads, Props -------------------------
    def month_series(sheet, header_row, max_col):
        teams, rows = {}, []
        for i, d in enumerate(wb.sheet(sheet, limit=700)):
            if i == header_row:
                for k, v in d.items():
                    t = text(v)
                    if t and k < max_col:
                        teams[k] = t
            dt = as_date(d.get(0))
            if teams and dt and 2015 < dt.year < 2040:
                vals = {t: num(d.get(k)) for k, t in teams.items() if num(d.get(k)) is not None}
                if vals:
                    rows.append((monthidx(dt), vals))
        return rows

    out['fte'] = month_series('FTE', 2, 21)
    out['leads_hist'] = month_series('Leads', 0, 19)
    out['props_hist'] = month_series('Props', 0, 18)

    # --- Locations: Hotel-Stammdaten -------------------------------------
    hotels = {}
    for i, d in enumerate(wb.sheet('Locations', limit=2000)):
        wid = num(d.get(0))
        if not wid or not (9000 < wid < 100000):
            continue
        name = text(d.get(2))
        if not name:
            continue
        seit = as_date(d.get(5))
        hotels[int(wid)] = dict(
            name=name,
            land=text(d.get(1)) or '',
            exkl=1 if text(d.get(7)) else 0,
            seit=monthidx(seit) if seit else None,
        )
    out['hotels'] = hotels

    # --- Goals: Annual Planning je Monat und Team ------------------------
    grid = list(wb.sheet('Goals', limit=44))
    ap = []
    for start, y0 in ((15, 2024), (52, 2025), (89, 2026)):
        for r in grid[3:35]:
            code = text(r.get(0))
            if not code or len(code) not in (2, 4) or code in ('FU', 'TU'):
                continue
            for off in range(12):
                v = num(r.get(start + off))
                if v is None:
                    continue
                mo, yr = 7 + off, y0
                if mo > 12:
                    mo, yr = mo - 12, y0 + 1
                ap.append((monthidx(datetime.date(yr, mo, 1)), code, v))
    out['ap'] = ap
    out['buq_fixed'], out['buq_year'] = read_fixed_buq(wb)
    return out


MONATE = {'januar': 1, 'februar': 2, 'märz': 3, 'maerz': 3, 'april': 4, 'mai': 5,
          'juni': 6, 'juli': 7, 'august': 8, 'september': 9, 'oktober': 10,
          'november': 11, 'dezember': 12}


def read_fixed_buq(wb):
    """Fest eingetragene Buchungsquoten aus der Tabelle "BuQ (Bu/Leads) JJJJ/JJ"
    im Blatt AP. Zellen mit Formel rechnet das Cockpit selbst; nur Werte, die
    von Hand eingetragen sind (z. B. FUNL, FUNW), ersetzen die Vorjahresquote.
    Liefert ({(Team, Kalendermonat): Quote}, Startjahr des AP-Zeitraums)."""
    try:
        grid = {}
        for r, c, v, is_f in wb.cells('AP', 80):
            grid[(r, c)] = (v, is_f)
    except SystemExit:
        return {}, None
    head = next(((r, c, v) for (r, c), (v, _) in grid.items()
                 if isinstance(v, str) and v.strip().lower().startswith('buq')), None)
    if not head:
        return {}, None
    hr, hc, label = head
    import re
    m = re.search(r'(20\d\d)\s*/\s*\d\d', label)
    if not m:
        return {}, None
    ap_year = int(m.group(1)) + 1                  # Quote 2025/26 -> AP 2026/27
    teams = {c: text(v) for (r, c), (v, _) in grid.items()
             if r == hr and c > hc and text(v)}
    fixed = {}
    for rr in range(hr + 1, hr + 13):
        mon = MONATE.get((text(grid.get((rr, hc), (None,))[0]) or '').lower())
        if not mon:
            continue
        for c, team in teams.items():
            v, is_f = grid.get((rr, c), (None, True))
            if not is_f and isinstance(v, (int, float)) and v > 0:
                fixed[(team, mon)] = float(v)
    return fixed, ap_year


# -------------------------------------------------------------- Combit
def csv_encoding(path):
    """UTF-8 (mit oder ohne BOM), sonst Windows-1252 - so speichern Excel und Combit."""
    try:
        with open(path, encoding='utf-8-sig') as f:
            for _ in f:
                pass
        return 'utf-8-sig'
    except UnicodeDecodeError:
        return 'cp1252'

def combit_rows(path):
    """Zeilen des Combit-Exports als dict {Spaltenname: Wert}."""
    ext = os.path.splitext(path)[1].lower()
    if ext == '.csv':
        with open(path, encoding=csv_encoding(path), newline='') as f:
            sample = f.read(8192)
            f.seek(0)
            delim = ';' if sample.count(';') >= sample.count(',') else ','
            for row in csv.DictReader(f, delimiter=delim):
                yield row
    else:
        wb = Workbook(path)
        it = wb.sheet(wb.names[0])
        header = None
        for d in it:
            if header is None:
                header = {k: (text(v) or f'col{k}') for k, v in d.items()}
                continue
            yield {name: d.get(k) for k, name in header.items()}


def read_combit(path):
    """Anfragen und Angebote, tagesgenau, je Team und je Hotel.

    Anfragen  Belegart 'Anfrage', dedupliziert je Ansprechpartner x Team x
              AP-Jahr; der frueheste Eintrag zaehlt. Das entspricht der
              Spalte 'einfach' in C_AP (unique Leads).
              Je Hotel zaehlt eine Anfrage bei JEDEM angefragten Hotel, einmal
              je Ansprechpartner x Team x AP-Jahr x Hotel. Die Summe ueber die
              Hotels ist deshalb groesser als die Zahl der unique Leads.
    Angebote  Belegart 'Angebot', datiert auf das Versanddatum, ersatzweise
              auf das Anfragedatum.
    Alle      jede Zeile mit Belegart 'Anfrage', ohne Deduplizierung, datiert
              auf ihr Anfragedatum — Basis der Buchungsquote (alle Anfragen).
    Web       Quelle 'website' oder Erfassung ueber das Webformular.
    Die Spalte AnsprechpartnerMail dient nur der Deduplizierung und wird
    nicht ausgegeben.
    """
    anfragen, angebote = {}, []
    hotel_anfragen = {}          # (mail, team, AP, WebID) -> (Datum, web)
    alle = []                    # jede Anfrage-Zeile, ohne Deduplizierung
    for row in combit_rows(path):
        team = text(row.get('Team'))
        herk = norm_country(text(row.get('KundenHerkunft')))
        ziel = norm_country(text(row.get('Destination')))
        a = as_date(row.get('Anfragedatum'))
        if a is None:
            continue
        wid = numlike(row.get('WebID'))
        wid = int(wid) if wid and 9000 < wid < 100000 else None
        web = (str(row.get('Quelle') or '').strip().lower() == 'website'
               or str(row.get('ErfassungsBenutzer') or '').strip() == 'Webformular_Import')
        art = (text(row.get('Belegart')) or '')
        if art == 'Anfrage':
            alle.append((a, team, wid, herk, ziel))
            key = (str(row.get('AnsprechpartnerMail') or '').strip().lower(),
                   team, str(row.get('AP') or '').strip())
            prev = anfragen.get(key)
            if prev is None or a < prev[0]:
                anfragen[key] = (a, team, web, wid, herk, ziel)
            if wid:
                hk = key + (wid,)
                hp = hotel_anfragen.get(hk)
                if hp is None or a < hp[0]:
                    hotel_anfragen[hk] = (a, web)
        elif art == 'Angebot':
            o = as_date(row.get('AngebotVersandtDatum'))
            if o is None or o.year < 2010:
                o = a
            angebote.append((o, team, wid, herk, ziel))

    # (Tag, Team, Herkunft, Reiseland) -> leads, web, offers, alle Anfragen
    by_team = defaultdict(lambda: [0, 0, 0, 0])
    by_hotel = defaultdict(lambda: [0, 0, 0, 0])    # (Tag, Hotel, Team) -> leads, web, offers, alle
    for a, team, web, wid, herk, ziel in anfragen.values():
        if team:
            r = by_team[(dayidx(a), team, herk, ziel)]
            r[0] += 1
            r[1] += 1 if web else 0
    for (_, team, _, wid), (a, web) in hotel_anfragen.items():
        r = by_hotel[(dayidx(a), wid, team)]; r[0] += 1; r[1] += 1 if web else 0
    for o, team, wid, herk, ziel in angebote:
        if team:
            by_team[(dayidx(o), team, herk, ziel)][2] += 1
        if wid:
            by_hotel[(dayidx(o), wid, team)][2] += 1
    # Alle Anfragen (ohne Unique-Regel) für die Buchungsquote je Team und Hotel
    for a, team, wid, herk, ziel in alle:
        if team:
            by_team[(dayidx(a), team, herk, ziel)][3] += 1
        if wid:
            by_hotel[(dayidx(a), wid, team)][3] += 1
    return by_team, by_hotel


# ------------------------------------------------------------ AP-Leads
def read_ap_leads(path):
    """Optionale Plandatei fuer Leads: Team;Monat;Leads.

    Ueberschreibt einzelne berechnete AP-Anfragen (compute_ap_leads), z. B.
    fuer manuell gesetzte Ziele. Monat als 2026-08 oder 08.2026.
    """
    out = []
    if not path or not os.path.exists(path):
        return out
    with open(path, encoding=csv_encoding(path), newline='') as f:
        sample = f.read(4096); f.seek(0)
        delim = ';' if sample.count(';') >= sample.count(',') else ','
        for row in csv.DictReader(f, delimiter=delim):
            team = text(row.get('Team'))
            mon = (row.get('Monat') or '').strip()
            val = numlike(row.get('Leads'))
            if not team or not mon or val is None:
                continue
            try:
                if '.' in mon:
                    mm, yy = mon.split('.')
                else:
                    yy, mm = mon.split('-')
                out.append((monthidx(datetime.date(int(yy), int(mm), 1)), team, val))
            except ValueError:
                continue
    return out


# ------------------------------------------------------------ AP-Anfragen
# Wie im Blatt AP der Mappe: Die Anfragen für einen Monat ergeben sich aus
# den geplanten Teams des FOLGEMONATS und der Buchungsquote dieses
# Folgemonats im Vorjahr — Anfragen kommen rund einen Monat vor der Buchung.
#
#   BuQ(M)       = Teams im Monat M des Vorjahres
#                  / unique Anfragen ("einfach") im Monat M-1 des Vorjahres
#   AP-Anfragen(m) = AP-Teams(m+1) / BuQ(m+1)
#
# Juni: Folgemonat ist der Juli desselben AP-Zeitraums. Ohne geplante oder
# stattgefundene Buchungen gibt es keine AP-Anfragen (0).
#
# Gezaehlt werden immer ganze Monate (alle Tage). Die Formel im Blatt AP
# vergleicht das Anfragedatum (mit Uhrzeit) mit "<=31.MM.JJJJ" und verliert
# dadurch die Anfragen vom letzten Tag des Monats. Nur zum Abgleich mit einer
# noch nicht korrigierten Mappe: AP_LEADS_WIE_EXCEL=1 bildet das nach.
AP_LEADS_WIE_EXCEL = os.environ.get('AP_LEADS_WIE_EXCEL', '0').strip() in ('1', 'ja', 'true')


def compute_ap_leads(wbd, by_team, cut, excel_compat=AP_LEADS_WIE_EXCEL):
    teams_m = defaultdict(float)                     # (Team, Monat) -> Teams
    for b in wbd['bookings']:
        if b['team']:
            teams_m[(b['team'], monthidx(DAY0 + datetime.timedelta(days=b['d'])))] += b['teams'] or 0
    leads_m = defaultdict(float)                     # (Team, Monat) -> unique Anfragen
    for (d, team, _h, _z), v in by_team.items():
        day = DAY0 + datetime.timedelta(days=d)
        if excel_compat:
            nxt = day + datetime.timedelta(days=1)
            if nxt.day == 1:                         # letzter Tag des Monats fehlt im Blatt AP
                continue
        leads_m[(team, monthidx(day))] += v[0]
    for mi, vals in wbd['leads_hist']:               # vor 2024 nur monatlich aus dem Blatt Leads
        if mi < cut:
            for team, v in vals.items():
                leads_m[(team, mi)] += v or 0
    ap = {(c, mi): v for mi, c, v in wbd['ap']}
    fixed, fixed_year = wbd.get('buq_fixed') or {}, wbd.get('buq_year')
    out = []
    for (team, m), _ in sorted(ap.items(), key=lambda x: (x[0][1], x[0][0])):
        nxt = m + 1 if m % 12 != 5 else m - 11       # Juni -> Juli desselben AP-Zeitraums
        plan = ap.get((team, nxt), 0) or 0
        ap_start = MONTH0[0] + (m - 6) // 12          # Startjahr des AP-Zeitraums (Juli)
        buq = fixed.get((team, nxt % 12 + 1)) if ap_start == fixed_year else None
        if buq:                                      # Quote im Blatt AP von Hand gesetzt
            val = plan / buq if plan > 0 else 0.0
        else:
            t_vj = teams_m.get((team, nxt - 12), 0)
            l_vj = leads_m.get((team, nxt - 13), 0)
            val = plan * l_vj / t_vj if plan > 0 and t_vj > 0 and l_vj > 0 else 0.0
        out.append((m, team, val))
    return out


# --------------------------------------------------------------- Aufbau
def build(xl, combit_path, outpath, ap_leads_path=None):
    wbd = read_workbook(xl)
    by_team, by_hotel = read_combit(combit_path)

    teams = sorted({b['team'] for b in wbd['bookings'] if b['team']} |
                   {k[1] for k in by_team})
    tix = {t: i for i, t in enumerate(teams)}

    # Hotels: Stammliste plus alles, was in Buchungen oder Combit auftaucht
    master = wbd['hotels']
    seen = {b['hotel'] for b in wbd['bookings'] if b['hotel']} | {h for (_, h, _) in by_hotel}
    fallback = {}
    for b in wbd['bookings']:
        if b['hotel'] and b['hotel'] not in master and b['name']:
            fallback.setdefault(b['hotel'], dict(name=b['name'], land=b['land'] or '',
                                                 exkl=0, seit=None))
    hid = sorted(set(master) | seen)
    hotels, hix = [], {}
    for i, h in enumerate(hid):
        m = master.get(h) or fallback.get(h) or dict(name=f'WebID {h}', land='', exkl=0, seit=None)
        hix[h] = i
        hotels.append([h, m['name'], m['land'], m['exkl'], m['seit']])

    sports = sorted({b['sport'] for b in wbd['bookings'] if b['sport']})
    # Nach Haeufigkeit sortieren - D, A, CH stehen sonst zwischen Codes,
    # die ein- oder zweimal vorkommen.
    def by_volume(getter, extra):
        cnt = defaultdict(int)
        for b in wbd['bookings']:
            v = getter(b)
            if v:
                cnt[v] += 1
        for v in extra:
            cnt.setdefault(v, 0)
        return sorted(cnt, key=lambda v: (-cnt[v], v))

    regs = by_volume(lambda b: b['reg'], {k[2] for k in by_team if k[2]})
    dests = by_volume(lambda b: b['land'], {k[3] for k in by_team if k[3]})
    six = {v: i for i, v in enumerate(sports)}
    rix = {v: i for i, v in enumerate(regs)}
    dix = {v: i for i, v in enumerate(dests)}

    # Regionen gehoeren immer zu einem Herkunftsland - das Kuerzel allein
    # ist nicht eindeutig (BE ist Berlin und Bern).
    rcount = defaultdict(int)
    for b in wbd['bookings']:
        if b['region'] and b['reg']:
            rcount[(b['reg'], b['region'])] += 1
    keys = sorted(rcount, key=lambda k: (regs.index(k[0]) if k[0] in regs else 99,
                                         -rcount[k], k[1]))
    regions = [[land, code, region_label(land, code)] for land, code in keys]
    gix = {k: i for i, k in enumerate(keys)}

    bookings = [[b['d'], tix.get(b['team'], -1), six.get(b['sport'], -1),
                 rix.get(b['reg'], -1), hix.get(b['hotel'], -1),
                 r2(b['vk']), r2(b['ek']), r2(b['marg']),
                 r2(b['teams'], 2), r2(b['pax'], 0), r2(b['nights'], 0),
                 gix.get((b['reg'], b['region']), -1), dix.get(b['land'], -1)]
                for b in wbd['bookings']]

    lead_team = sorted([[d, tix[t], rix.get(h, -1), dix.get(z, -1), *v]
                        for (d, t, h, z), v in by_team.items() if t in tix])
    # leadDaily: [Tag, Team, Herkunft, Reiseland, Anfragen unique, Web, Angebote, alle Anfragen]
    # leadHotel: [Tag, Hotel, Team (-1 = ohne Team), Anfragen, Web, Angebote, alle Anfragen]
    lead_hotel = sorted([[d, hix[h], tix.get(t, -1), *v] for (d, h, t), v in by_hotel.items() if h in hix])

    def mrows(series, cut=None):
        out = []
        for mi, vals in series:
            if cut is not None and mi >= cut:
                continue
            for t, v in vals.items():
                if t in tix:
                    out.append([mi, tix[t], r2(v, 3)])
        return out

    cut = monthidx(datetime.date(*LEAD_CUT, 1))
    # AP-Anfragen: berechnet wie im Blatt AP; eine AP_Leads.csv ersetzt einzelne Werte
    apl = {(mi, c): v for mi, c, v in compute_ap_leads(wbd, by_team, cut)}
    csv_rows = read_ap_leads(ap_leads_path)
    for mi, c, v in csv_rows:
        apl[(mi, c)] = v
    ap_leads_rows = [[mi, tix[c], r2(v, 2)] for (mi, c), v in sorted(apl.items()) if c in tix]
    data = dict(
        meta=dict(
            generated=datetime.datetime.now().isoformat(timespec='seconds'),
            day0=DAY0.isoformat(),
            month0=f'{MONTH0[0]}-{MONTH0[1]:02d}',
            leadCutMonth=cut,
            sources=dict(workbook=os.path.basename(xl), combit=os.path.basename(combit_path)),
        ),
        teams=teams, sports=sports, regs=regs, dests=dests,
        regions=regions, countries=COUNTRY, hotels=hotels,
        bookings=bookings,
        leadDaily=lead_team,
        leadHotel=lead_hotel,
        leadMonthly=mrows(wbd['leads_hist'], cut),
        propMonthly=mrows(wbd['props_hist'], cut),
        fte=mrows(wbd['fte']),
        ap=[[mi, tix[c], r2(v, 3)] for mi, c, v in wbd['ap'] if c in tix],
        apLeads=ap_leads_rows,
    )
    os.makedirs(os.path.dirname(os.path.abspath(outpath)), exist_ok=True)
    tmp = outpath + '.tmp'
    with open(tmp, 'w', encoding='utf-8') as f:
        json.dump(data, f, separators=(',', ':'), ensure_ascii=False)
    os.replace(tmp, outpath)

    print(f'Teams            {len(teams)}: {", ".join(teams)}')
    print(f'Hotels           {len(hotels)} (Stammliste {len(master)}, ohne Stammsatz {len(fallback)})')
    print(f'Herkunftslaender {len(regs)} | Regionen {len(regions)} | Reiselaender {len(dests)}')
    print(f'Buchungen        {len(bookings)}')
    print(f'Leads je Team    {len(lead_team)} Zeilen')
    print(f'Leads je Hotel   {len(lead_hotel)} Zeilen')
    print(f'FTE              {len(data["fte"])} Zeilen')
    print(f'Annual Planning  {len(data["ap"])} Zeilen')
    print(f'AP-Anfragen      {len(data["apLeads"])} Zeilen, berechnet aus AP-Teams und Vorjahresquote'
          + (' (wie Blatt AP, ohne letzten Monatstag)' if AP_LEADS_WIE_EXCEL else ' (ganze Monate)')
          + (f', {len(csv_rows)} Werte aus AP_Leads.csv' if csv_rows else ''))
    if wbd.get('buq_fixed'):
        fx = sorted({t for t, _ in wbd['buq_fixed']})
        print(f'Feste BuQ        AP {wbd["buq_year"]}/{(wbd["buq_year"] + 1) % 100:02d} aus Blatt AP: {", ".join(fx)}')
    print(f'-> {outpath}  {os.path.getsize(outpath)/1e6:.2f} MB')


if __name__ == '__main__':
    args = sys.argv[1:]
    if len(args) < 2:
        raise SystemExit(__doc__.strip())
    build(args[0], args[1], args[2] if len(args) > 2 else 'data.json',
          args[3] if len(args) > 3 else None)
